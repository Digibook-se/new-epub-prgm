# converter.py
import os
import base64
import traceback
from ebooklib import epub
from ebooklib.epub import EpubHtml, EpubNav, EpubNcx, Link
from ebooklib.utils import create_pagebreak
from docx import Document
from docx.opc.exceptions import PackageNotFoundError
import PyPDF2

from functions.create import createEbook
from functions.createHelpers import (
    getCoverImg, getOpeningElementTag, getClosingElementTag,
    getStyle, removeFile
)
from functions.opfAndTocModifyer import setCorrectMetatags, modifyOpf, modifyToc
from structuralComponents.structure import getCoverPageXml
from miniComponents.copyright import getCopyrightPage
from classes.BookToBuild import BookToBuild
from Styles import base
from colorama import Fore


def parse_word(path: str):
    """Läser .docx och returnerar listan bookToBuild.Content-format."""
    try:
        doc = Document(path)
    except PackageNotFoundError:
        raise ValueError(f"Filen är inte ett giltigt .docx-paket:\n{path}")
    except Exception as e:
        raise ValueError(f"Misslyckades läsa Word-filen:\n{e}")

    if not doc.paragraphs:
        raise ValueError("Word-dokumentet är tomt (inga stycken).")

    content = []
    uid = 1
    for para in doc.paragraphs:
        runs = []
        for run in para.runs:
            attrs = []
            if run.bold:      attrs.append("bold")
            if run.italic:    attrs.append("italic")
            if run.underline: attrs.append("underline")
            runs.append({
                "Text": run.text or "",
                "Type": "",
                "Attributes": attrs
            })
        content.append({"Id": f"PAR_{uid}", "Paragraph": runs})
        uid += 1
    return content


def parse_pdf(path: str):
    """Läser PDF-sidor rad för rad och returnerar samma format."""
    try:
        reader = PyPDF2.PdfReader(path)
    except Exception as e:
        raise ValueError(f"Misslyckades läsa PDF-filen:\n{e}")

    content = []
    uid = 1
    for page in reader.pages:
        text = page.extract_text() or ""
        for line in text.splitlines():
            content.append({
                "Id": f"PAR_{uid}",
                "Paragraph": [{"Text": line, "Type": "", "Attributes": []}]
            })
            uid += 1
    if not content:
        raise ValueError("PDF-dokumentet är tomt (ingen text kunde extraheras).")
    return content


def convert_file(metadata: dict, include_toc: bool):
    """
    Gör hela kedjan: parsar Word/PDF, packar metadata + content i newBook,
    anropar createEbook(), tar hand om alla hjälpfunktioner.
    """
    path = metadata.get("path", "")
    ext = os.path.splitext(path)[1].lower()

    # 1) Parsning av indatafil
    if ext == ".docx":
        content = parse_word(path)
    elif ext == ".pdf":
        content = parse_pdf(path)
    else:
        raise ValueError(f"Obehandlat filformat: {ext}")

    # 2) Sätt ihop newBook-dict som createEbook förväntar sig
    newBook = {
        "title":         metadata.get("title", ""),
        "author":        metadata.get("author", ""),
        "photographer":  metadata.get("photographer", ""),
        "ISBN":          metadata.get("ISBN", ""),
        "categories":    metadata.get("categories", []),
        "description":   metadata.get("description", ""),
        "content":       content,
        "toc":           include_toc
    }

    # 3) Skapa EPUB via din befintliga funktion
    try:
        createEbook(newBook)
    except Exception as e:
        # Lägg gärna logga stack-trace till en fil om du vill
        traceback.print_exc()
        raise
